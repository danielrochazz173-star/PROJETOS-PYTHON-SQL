"""
Relatório de Vendas por Departamento e Mês
Permite selecionar departamento e múltiplos meses
Gera planilha com: Vendedor, QtVendido, Produto, CodProduto, EmbMaster, CodCli, Nome Cliente
"""

import tkinter as tk
from tkinter import ttk, messagebox
import oracledb
import pandas as pd
from datetime import datetime
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import threading
from calendar import monthrange

# Configuração do Oracle Instant Client
oracledb.init_oracle_client(lib_dir=r"C:\oracle\instantclient_23_9")

# =========================================================================
# CONFIGURAÇÕES
# =========================================================================
DB_HOST = os.getenv('DB_HOST', '10.0.0.10')
DB_PORT = int(os.getenv('DB_PORT', '1521'))
DB_SERVICE = os.getenv('DB_SERVICE', 'PROD')
DB_USER = os.getenv('DB_USER', 'powerbi')
DB_PASSWORD = os.getenv('DB_PASSWORD', '')

# =========================================================================
# FUNÇÕES DE BANCO
# =========================================================================
def get_db_connection():
    """Conecta ao banco Oracle"""
    try:
        return oracledb.connect(
            user=DB_USER,
            password=DB_PASSWORD,
            dsn=f"{DB_HOST}:{DB_PORT}/{DB_SERVICE}",
        )
    except oracledb.DatabaseError as e:
        return None

def buscar_departamentos(conn):
    """Busca lista de departamentos"""
    try:
        query = "SELECT CODEPTO, DESCRICAO FROM PCDEPTO ORDER BY DESCRICAO"
        df = pd.read_sql_query(query, conn)
        df.columns = [x.lower() for x in df.columns]
        return df
    except Exception as e:
        print(f"Erro ao buscar departamentos: {e}")
        return pd.DataFrame()

# =========================================================================
# CLASSE PRINCIPAL
# =========================================================================
class RelatorioVendasDeptoMes:
    def __init__(self, root):
        self.root = root
        self.root.title("📊 Relatório de Vendas por Departamento e Mês")
        self.root.geometry("600x700")
        
        # Variáveis
        self.conn = None
        self.df_departamentos = None
        self.meses_selecionados = {}  # {ano_mes: (ano, mes)}
        
        # Criar interface
        self.criar_interface()
        
        # Conectar ao banco em thread separada
        self.conectar_banco()
    
    def criar_interface(self):
        """Cria a interface gráfica"""
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Título
        titulo = tk.Label(main_frame, text="📊 Relatório de Vendas por Departamento e Mês", 
                          font=("Arial", 14, "bold"))
        titulo.pack(pady=(0, 20))
        
        # Departamento
        frame_depto = ttk.Frame(main_frame)
        frame_depto.pack(fill=tk.X, pady=10)
        ttk.Label(frame_depto, text="Departamento:", width=20).pack(side=tk.LEFT)
        self.combo_depto = ttk.Combobox(frame_depto, width=40, state="readonly")
        self.combo_depto.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        
        # Vendedores (CODUSUR/RCA)
        frame_vendedores = ttk.Frame(main_frame)
        frame_vendedores.pack(fill=tk.X, pady=10)
        ttk.Label(frame_vendedores, text="Vendedores (CODUSUR):", width=20).pack(side=tk.LEFT)
        self.entry_vendedores = ttk.Entry(frame_vendedores, width=40)
        self.entry_vendedores.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Label(frame_vendedores, text="(separados por vírgula, ex: 1,2,3)").pack(side=tk.LEFT, padx=5)
        
        # Ano
        frame_ano = ttk.Frame(main_frame)
        frame_ano.pack(fill=tk.X, pady=10)
        ttk.Label(frame_ano, text="Ano:", width=20).pack(side=tk.LEFT)
        self.combo_ano = ttk.Combobox(frame_ano, width=40, state="readonly", values=["2025", "2026"])
        self.combo_ano.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        # Definir ano padrão (ano atual ou 2025 se ainda não for 2025)
        ano_atual = datetime.now().year
        if ano_atual >= 2026:
            self.combo_ano.current(1)  # 2026
        else:
            self.combo_ano.current(0)  # 2025
        self.combo_ano.bind("<<ComboboxSelected>>", lambda e: self.atualizar_checkboxes_meses())
        
        # Frame para meses
        frame_meses_label = ttk.Frame(main_frame)
        frame_meses_label.pack(fill=tk.X, pady=(20, 5))
        ttk.Label(frame_meses_label, text="Selecione os meses:").pack(side=tk.LEFT)
        
        # Frame com scroll para meses
        frame_meses_container = ttk.Frame(main_frame)
        frame_meses_container.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # Canvas e scrollbar para meses
        canvas = tk.Canvas(frame_meses_container, height=300)
        scrollbar = ttk.Scrollbar(frame_meses_container, orient="vertical", command=canvas.yview)
        scrollable_frame = ttk.Frame(canvas)
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # Criar checkboxes de meses
        self.frame_meses = scrollable_frame
        self.criar_checkboxes_meses()
        
        # Botão de gerar Excel
        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=(20, 0))
        
        self.btn_gerar = ttk.Button(btn_frame, text="📊 Gerar Excel", 
                                   command=self.gerar_excel, state=tk.DISABLED,
                                   style="Accent.TButton")
        self.btn_gerar.pack(side=tk.LEFT, padx=(0, 10))
        
        self.status_label = ttk.Label(btn_frame, text="Conectando ao banco...")
        self.status_label.pack(side=tk.LEFT)
    
    def criar_checkboxes_meses(self):
        """Cria checkboxes para seleção de meses"""
        # Limpar frame
        for widget in self.frame_meses.winfo_children():
            widget.destroy()
        
        # Obter ano selecionado
        ano_selecionado = int(self.combo_ano.get()) if self.combo_ano.get() else datetime.now().year
        
        # Criar checkboxes para os 12 meses do ano selecionado
        self.checkboxes_meses = {}
        
        meses_pt = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
                   'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
        
        for mes in range(1, 13):
            nome_mes = f"{meses_pt[mes-1]}/{ano_selecionado}"
            chave = f"{ano_selecionado}_{mes:02d}"
            
            # Criar checkbox
            var = tk.BooleanVar()
            checkbox = ttk.Checkbutton(
                self.frame_meses, 
                text=nome_mes,
                variable=var
            )
            checkbox.pack(anchor=tk.W, padx=10, pady=2)
            
            self.checkboxes_meses[chave] = {
                'var': var,
                'ano': ano_selecionado,
                'mes': mes,
                'nome': nome_mes
            }
    
    def atualizar_checkboxes_meses(self):
        """Atualiza os checkboxes quando o ano é alterado"""
        self.criar_checkboxes_meses()
    
    def conectar_banco(self):
        """Conecta ao banco e carrega dados em thread separada"""
        def carregar():
            try:
                self.status_label.config(text="Conectando ao banco...")
                self.conn = get_db_connection()
                
                if not self.conn:
                    self.status_label.config(text="❌ Erro ao conectar")
                    messagebox.showerror("Erro", "Não foi possível conectar ao banco de dados")
                    return
                
                self.status_label.config(text="Carregando departamentos...")
                self.df_departamentos = buscar_departamentos(self.conn)
                
                if not self.df_departamentos.empty:
                    deptos = ["Todos os Departamentos"] + [f"{int(row['codepto'])} - {row['descricao']}" 
                             for _, row in self.df_departamentos.iterrows()]
                    self.combo_depto['values'] = deptos
                    if deptos:
                        self.combo_depto.current(0)
                else:
                    self.combo_depto['values'] = ["Todos os Departamentos"]
                
                self.status_label.config(text="✅ Pronto para gerar relatório")
                self.btn_gerar.config(state=tk.NORMAL)
                
            except Exception as e:
                self.status_label.config(text=f"❌ Erro: {str(e)}")
                messagebox.showerror("Erro", f"Erro ao carregar dados: {e}")
        
        threading.Thread(target=carregar, daemon=True).start()
    
    def obter_meses_selecionados(self):
        """Retorna lista de (ano, mes) dos meses selecionados"""
        meses = []
        for chave, dados in self.checkboxes_meses.items():
            if dados['var'].get():
                meses.append((dados['ano'], dados['mes']))
        return sorted(meses, key=lambda x: (x[0], x[1]))
    
    def montar_query(self, ano, mes):
        """Monta a query SQL para um mês específico"""
        # Obter primeiro e último dia do mês
        primeiro_dia = f"01/{mes:02d}/{ano}"
        ultimo_dia = f"{monthrange(ano, mes)[1]:02d}/{mes:02d}/{ano}"
        
        # Obter código do departamento
        depto_sel = self.combo_depto.get()
        if not depto_sel:
            raise ValueError("Selecione um departamento")
        
        # Filtro de departamento
        filtro_depto = ""
        if depto_sel != "Todos os Departamentos":
            coddepto = depto_sel.split(" - ")[0]
            filtro_depto = f"AND P.CODEPTO = {coddepto}"
        
        # Filtro de vendedores (CODUSUR)
        filtro_vendedores = ""
        vendedores_str = self.entry_vendedores.get().strip()
        if vendedores_str:
            codusur_list = []
            for cod in vendedores_str.split(','):
                cod_limpo = cod.strip()
                if cod_limpo.isdigit():
                    codusur_list.append(int(cod_limpo))
            
            if codusur_list:
                # Dividir em chunks de 1000 se necessário
                if len(codusur_list) <= 1000:
                    filtro_vendedores = f"AND C.CODUSUR IN ({', '.join(map(str, codusur_list))})"
                else:
                    # Múltiplos chunks - usar OR
                    chunks = [codusur_list[i:i+1000] for i in range(0, len(codusur_list), 1000)]
                    condicoes = []
                    for chunk in chunks:
                        condicoes.append(f"C.CODUSUR IN ({', '.join(map(str, chunk))})")
                    filtro_vendedores = f"AND ({' OR '.join(condicoes)})"
        
        query = f"""
        SELECT 
            U.NOME AS NOME_VENDEDOR,
            I.QT AS QTVENDIDO,
            P.DESCRICAO AS PRODUTO,
            I.CODPROD AS CODPRODUTO,
            NVL(P.QTUNITCX, 1) AS EMBMASTER,
            CL.CODCLI,
            CL.CLIENTE AS NOME_CLIENTE
        FROM PCPEDC C
        JOIN PCPEDI I ON C.NUMPED = I.NUMPED
        JOIN PCCLIENT CL ON C.CODCLI = CL.CODCLI
        JOIN PCPRODUT P ON I.CODPROD = P.CODPROD
        LEFT JOIN PCUSUARI U ON C.CODUSUR = U.CODUSUR
        WHERE C.DATA BETWEEN TO_DATE('{primeiro_dia}', 'DD/MM/YYYY') 
            AND TO_DATE('{ultimo_dia}', 'DD/MM/YYYY')
            {filtro_depto}
            {filtro_vendedores}
            AND C.CODFILIAL IN ('1', '98')
            AND C.POSICAO = 'F'
            AND C.DTCANCEL IS NULL
            AND C.CONDVENDA NOT IN (4, 8, 10, 13, 20, 98, 99)
            AND NVL(I.BONIFIC, 'N') = 'N'
        ORDER BY U.NOME, CL.CLIENTE, P.DESCRICAO
        """
        
        return query
    
    def gerar_excel(self):
        """Gera o arquivo Excel"""
        if not self.conn:
            messagebox.showerror("Erro", "Não há conexão com o banco de dados")
            return
        
        # Verificar se há departamento selecionado
        depto_sel = self.combo_depto.get()
        if not depto_sel:
            messagebox.showerror("Erro", "Selecione um departamento")
            return
        
        # Verificar se há meses selecionados
        meses = self.obter_meses_selecionados()
        if not meses:
            messagebox.showerror("Erro", "Selecione pelo menos um mês")
            return
        
        def executar():
            try:
                self.status_label.config(text="Buscando dados...")
                
                # Buscar dados de todos os meses selecionados
                dfs_mes = []
                total_registros = 0
                
                for ano, mes in meses:
                    meses_pt = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
                               'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
                    self.status_label.config(
                        text=f"Processando {meses_pt[mes-1]}/{ano}..."
                    )
                    query = self.montar_query(ano, mes)
                    df = pd.read_sql_query(query, self.conn)
                    
                    if not df.empty:
                        # Adicionar coluna de mês/ano
                        df['MES_ANO'] = datetime(ano, mes, 1).strftime("%m/%Y")
                        dfs_mes.append(df)
                        total_registros += len(df)
                
                if not dfs_mes:
                    messagebox.showwarning("Aviso", "Nenhum registro encontrado com os filtros especificados")
                    self.status_label.config(text="❌ Nenhum registro encontrado")
                    return
                
                # Concatenar todos os dataframes
                df_final = pd.concat(dfs_mes, ignore_index=True)
                
                # Reordenar colunas: MES_ANO primeiro, depois as outras
                colunas = ['MES_ANO', 'NOME_VENDEDOR', 'QTVENDIDO', 'PRODUTO', 
                          'CODPRODUTO', 'EMBMASTER', 'CODCLI', 'NOME_CLIENTE']
                df_final = df_final[colunas]
                
                self.status_label.config(text=f"Gerando Excel com {total_registros} registros...")
                
                # Nome do arquivo
                if depto_sel == "Todos os Departamentos":
                    depto_nome = "Todos_Departamentos"
                else:
                    depto_nome = depto_sel.split(" - ")[1] if " - " in depto_sel else depto_sel
                nome_arquivo = f"vendas_depto_{depto_nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
                caminho_arquivo = Path(nome_arquivo).resolve()
                
                # Gerar Excel formatado
                self.formatar_excel(caminho_arquivo, df_final)
                
                self.status_label.config(text="✅ Excel gerado com sucesso!")
                messagebox.showinfo("Sucesso", f"Excel gerado:\n{caminho_arquivo}\n\nTotal de registros: {total_registros}")
                
                # Abrir Excel
                try:
                    os.startfile(str(caminho_arquivo))
                except:
                    pass
                    
            except ValueError as e:
                messagebox.showerror("Erro", str(e))
                self.status_label.config(text="❌ Erro na validação")
            except Exception as e:
                messagebox.showerror("Erro", f"Erro ao gerar Excel:\n{e}")
                self.status_label.config(text="❌ Erro ao gerar")
        
        threading.Thread(target=executar, daemon=True).start()
    
    def formatar_excel(self, arquivo, df):
        """Formata o arquivo Excel"""
        with pd.ExcelWriter(arquivo, engine='openpyxl') as writer:
            # Aba de dados
            df.to_excel(writer, sheet_name='Vendas', index=False)
            ws = writer.sheets['Vendas']
            
            # Estilos
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin'),
            )
            
            # Formatar cabeçalho
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            # Ajustar larguras das colunas
            col_widths = {
                'A': 12,  # MES_ANO
                'B': 25,  # NOME_VENDEDOR
                'C': 12,  # QTVENDIDO
                'D': 40,  # PRODUTO
                'E': 12,  # CODPRODUTO
                'F': 12,  # EMBMASTER
                'G': 12,  # CODCLI
                'H': 40   # NOME_CLIENTE
            }
            for col, width in col_widths.items():
                ws.column_dimensions[col].width = width
            
            # Formatar dados
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(horizontal='left', vertical='center')
            
            # Formatar números
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                # QTVENDIDO (coluna C)
                if row[2].value is not None:
                    row[2].number_format = '#,##0.00'
                # EMBMASTER (coluna F)
                if row[5].value is not None:
                    row[5].number_format = '#,##0'
            
            # Congelar primeira linha
            ws.freeze_panes = 'A2'

# =========================================================================
# EXECUÇÃO
# =========================================================================
if __name__ == "__main__":
    root = tk.Tk()
    app = RelatorioVendasDeptoMes(root)
    root.mainloop()

# portfolio-commit-ready: relatorio_vendas_depto_mes
