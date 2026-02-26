"""
Dashboard Streamlit - Análise de Vendedores
Tabela completa com comparação mês atual vs anterior
"""

import streamlit as st
import pandas as pd
import oracledb
from datetime import datetime, timedelta
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configuração do Oracle Instant Client
oracledb.init_oracle_client(lib_dir=r"C:\oracle\instantclient_23_9")

# =========================================================================
# CONFIGURAÇÕES
# =========================================================================
DB_HOST = '10.0.0.10'
DB_PORT = 1521
DB_SERVICE = 'PROD'
DB_USER = 'powerbi'
DB_PASSWORD = 'cbjc4xp3nlq6'

# =========================================================================
# FUNÇÕES DE BANCO
# =========================================================================
@st.cache_resource
def get_db_connection():
    """Conecta ao banco Oracle"""
    try:
        return oracledb.connect(
            user=DB_USER,
            password=DB_PASSWORD,
            dsn=f"{DB_HOST}:{DB_PORT}/{DB_SERVICE}",
        )
    except Exception as e:
        st.error(f"Erro ao conectar ao banco: {e}")
        return None

def buscar_supervisores(_conn):
    """Busca lista de supervisores"""
    try:
        query = """
        SELECT DISTINCT CODSUPERVISOR, NOME
        FROM PCSUPERV
        ORDER BY CODSUPERVISOR
        """
        df = pd.read_sql_query(query, _conn)
        df.columns = [x.lower() for x in df.columns]
        return df
    except:
        return pd.DataFrame()

def buscar_dados_vendedores(_conn, mes, ano, dia_fim=None, supervisores_list=None):
    """
    Busca dados completos de vendedores para um mês específico
    Retorna: faturamento, pedidos (F+L), margem, positivação, mix, qtd NF
    """
    # Data início do mês
    data_inicio = datetime(ano, mes, 1)
    
    # Data fim (até o dia especificado ou fim do mês)
    if dia_fim:
        try:
            data_fim = datetime(ano, mes, dia_fim)
        except:
            data_fim = datetime(ano, mes, 1)
            if mes == 12:
                data_fim = datetime(ano + 1, 1, 1)
            else:
                data_fim = datetime(ano, mes + 1, 1)
            data_fim = data_fim - timedelta(days=1)
    else:
        if mes == 12:
            data_fim = datetime(ano + 1, 1, 1)
        else:
            data_fim = datetime(ano, mes + 1, 1)
        data_fim = data_fim - timedelta(days=1)
    
    data_inicio_str = data_inicio.strftime('%d/%m/%Y')
    data_fim_str = data_fim.strftime('%d/%m/%Y')
    
    # Filtro de supervisores
    filtro_supervisor = ""
    if supervisores_list:
        supervisores_str = ', '.join(map(str, supervisores_list))
        filtro_supervisor = f"AND U.CODSUPERVISOR IN ({supervisores_str})"
    
    query = f"""
    WITH CLIENTES_POSITIVADOS AS (
        SELECT DISTINCT
            C.CODUSUR,
            C.CODCLI
        FROM PCPEDC C 
        JOIN PCPEDI I ON C.NUMPED = I.NUMPED
        JOIN PCPRODUT P ON I.CODPROD = P.CODPROD
        JOIN PCUSUARI U ON C.CODUSUR = U.CODUSUR
        WHERE C.DATA >= TO_DATE('{data_inicio_str}', 'DD/MM/YYYY')
          AND C.DATA <= TO_DATE('{data_fim_str}', 'DD/MM/YYYY')
          AND C.CODFILIAL IN ('1', '98')
          AND C.POSICAO = 'F'
          AND C.DTCANCEL IS NULL
          AND C.CONDVENDA NOT IN (4, 8, 10, 13, 20, 98, 99)
          AND NVL(I.BONIFIC, 'N') = 'N'
          {filtro_supervisor}
    ),
    VENDAS_FATURADAS AS (
        SELECT 
            C.CODUSUR,
            U.NOME AS NOME_VENDEDOR,
            C.CODCLI,
            CL.CLIENTE AS NOME_CLIENTE,
            SUM(
                CASE 
                    WHEN NVL(I.BONIFIC, 'N') = 'N' THEN 
                        DECODE(C.CONDVENDA, 
                            5, 0, 6, 0, 11, 0, 12, 0, 
                            ROUND(NVL(I.QT, 0) * NVL(I.PVENDA, 0), 2)
                        )
                    ELSE 0 
                END
            ) AS VALOR_BRUTO,
            SUM(
                CASE 
                    WHEN NVL(I.BONIFIC, 'N') = 'N' THEN 
                        DECODE(C.CONDVENDA, 
                            5, 0, 6, 0, 11, 0, 12, 0, 
                            NVL(I.QT, 0) * NVL(I.VLCUSTOFIN, 0)
                        )
                    ELSE 0 
                END
            ) AS CUSTO_BRUTO,
            COUNT(DISTINCT I.CODPROD) AS QTD_PRODUTOS_DIFERENTES,
            COUNT(DISTINCT C.NUMNOTA) AS QTD_NF
        FROM PCPEDC C 
        JOIN PCPEDI I ON C.NUMPED = I.NUMPED
        JOIN PCPRODUT P ON I.CODPROD = P.CODPROD
        JOIN PCUSUARI U ON C.CODUSUR = U.CODUSUR
        JOIN PCCLIENT CL ON C.CODCLI = CL.CODCLI
        WHERE C.DATA >= TO_DATE('{data_inicio_str}', 'DD/MM/YYYY')
          AND C.DATA <= TO_DATE('{data_fim_str}', 'DD/MM/YYYY')
          AND C.CODFILIAL IN ('1', '98')
          AND C.POSICAO = 'F'
          AND C.DTCANCEL IS NULL
          AND C.CONDVENDA NOT IN (4, 8, 10, 13, 20, 98, 99)
          {filtro_supervisor}
        GROUP BY C.CODUSUR, U.NOME, C.CODCLI, CL.CLIENTE
    ),
    PEDIDOS_LIBERADOS AS (
        SELECT 
            C.CODUSUR,
            C.CODCLI,
            SUM(
                CASE 
                    WHEN NVL(I.BONIFIC, 'N') = 'N' THEN 
                        DECODE(C.CONDVENDA, 
                            5, 0, 6, 0, 11, 0, 12, 0, 
                            ROUND(NVL(I.QT, 0) * NVL(I.PVENDA, 0), 2)
                        )
                    ELSE 0 
                END
            ) AS VALOR_PEDIDOS
        FROM PCPEDC C 
        JOIN PCPEDI I ON C.NUMPED = I.NUMPED
        JOIN PCPRODUT P ON I.CODPROD = P.CODPROD
        JOIN PCUSUARI U ON C.CODUSUR = U.CODUSUR
        WHERE C.DATA >= TO_DATE('{data_inicio_str}', 'DD/MM/YYYY')
          AND C.DATA <= TO_DATE('{data_fim_str}', 'DD/MM/YYYY')
          AND C.CODFILIAL IN ('1', '98')
          AND C.POSICAO = 'L'
          AND C.DTCANCEL IS NULL
          AND C.CONDVENDA NOT IN (4, 8, 10, 13, 20, 98, 99)
          {filtro_supervisor}
        GROUP BY C.CODUSUR, C.CODCLI
    ),
    DEVOLUCOES AS (
        SELECT 
            D.CODUSUR,
            SUM(NVL(D.VLDEVOLUCAO, 0)) AS VALOR_DEVOLVIDO,
            SUM(NVL(D.VLCUSTOFIN, 0)) AS CUSTO_DEVOLVIDO
        FROM VIEW_DEVOL_RESUMO_FATURAMENTO D
        JOIN PCPRODUT P ON D.CODPROD = P.CODPROD
        JOIN PCUSUARI U ON D.CODUSUR = U.CODUSUR
        WHERE D.DTENT >= TO_DATE('{data_inicio_str}', 'DD/MM/YYYY')
          AND D.DTENT <= TO_DATE('{data_fim_str}', 'DD/MM/YYYY')
          {filtro_supervisor}
        GROUP BY D.CODUSUR
    ),
    POSITIVACAO_POR_VENDEDOR_CLIENTE AS (
        SELECT 
            CODUSUR,
            CODCLI,
            CASE WHEN COUNT(*) > 0 THEN 1 ELSE 0 END AS POSITIVACAO
        FROM CLIENTES_POSITIVADOS
        GROUP BY CODUSUR, CODCLI
    ),
    DEVOLUCOES_POR_CLIENTE AS (
        SELECT 
            D.CODUSUR,
            C.CODCLI,
            SUM(NVL(D.VLDEVOLUCAO, 0)) AS VALOR_DEVOLVIDO,
            SUM(NVL(D.VLCUSTOFIN, 0)) AS CUSTO_DEVOLVIDO
        FROM VIEW_DEVOL_RESUMO_FATURAMENTO D
        JOIN PCPRODUT P ON D.CODPROD = P.CODPROD
        JOIN PCUSUARI U ON D.CODUSUR = U.CODUSUR
        JOIN PCPEDC C ON D.NUMPED = C.NUMPED
        WHERE D.DTENT >= TO_DATE('{data_inicio_str}', 'DD/MM/YYYY')
          AND D.DTENT <= TO_DATE('{data_fim_str}', 'DD/MM/YYYY')
          {filtro_supervisor}
        GROUP BY D.CODUSUR, C.CODCLI
    )
    SELECT 
        V.CODUSUR AS RCA,
        V.NOME_VENDEDOR,
        V.CODCLI,
        V.NOME_CLIENTE,
        (NVL(V.VALOR_BRUTO, 0) - NVL(D.VALOR_DEVOLVIDO, 0)) AS VENDA_MES,
        (NVL(V.VALOR_BRUTO, 0) - NVL(D.VALOR_DEVOLVIDO, 0) + NVL(L.VALOR_PEDIDOS, 0)) AS VENDA_PEDIDO,
        (NVL(V.VALOR_BRUTO, 0) - NVL(D.VALOR_DEVOLVIDO, 0)) - (NVL(V.CUSTO_BRUTO, 0) - NVL(D.CUSTO_DEVOLVIDO, 0)) AS MARGEM_VALOR,
        NVL(P.POSITIVACAO, 0) AS POSITIVACAO,
        NVL(V.QTD_PRODUTOS_DIFERENTES, 0) AS MIX,
        NVL(V.QTD_NF, 0) AS QTD_NF
    FROM VENDAS_FATURADAS V
    LEFT JOIN DEVOLUCOES_POR_CLIENTE D ON V.CODUSUR = D.CODUSUR AND V.CODCLI = D.CODCLI
    LEFT JOIN PEDIDOS_LIBERADOS L ON V.CODUSUR = L.CODUSUR AND V.CODCLI = L.CODCLI
    LEFT JOIN POSITIVACAO_POR_VENDEDOR_CLIENTE P ON V.CODUSUR = P.CODUSUR AND V.CODCLI = P.CODCLI
    WHERE V.VALOR_BRUTO > 0
    ORDER BY V.CODUSUR, VENDA_MES DESC
    """
    
    try:
        df = pd.read_sql_query(query, _conn)
        df.columns = [x.lower() for x in df.columns]
        
        # Calcular margem percentual
        if not df.empty:
            df['margem_percent'] = df.apply(
                lambda row: 0.0 if row['venda_mes'] <= 0 
                else round(((row['margem_valor'] / row['venda_mes']) * 100), 2),
                axis=1
            )
        
        return df
    except Exception as e:
        st.error(f"Erro ao buscar dados: {e}")
        return pd.DataFrame()

# =========================================================================
# INTERFACE STREAMLIT
# =========================================================================
st.set_page_config(
    page_title="Dashboard Vendedores",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Título
st.title("📊 Dashboard de Análise de Vendedores")

# Sidebar - Filtros
st.sidebar.header("🔧 Filtros")

# Seleção de mês
meses = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
    5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
    9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro"
}

hoje = datetime.now()
mes_selecionado = st.sidebar.selectbox(
    "Mês",
    options=list(meses.keys()),
    format_func=lambda x: f"{meses[x]} {hoje.year}",
    index=hoje.month - 1
)

ano_selecionado = hoje.year

# Seleção de dia (até qual dia do mês)
dia_max = 31
if mes_selecionado in [4, 6, 9, 11]:
    dia_max = 30
elif mes_selecionado == 2:
    dia_max = 29 if (ano_selecionado % 4 == 0 and ano_selecionado % 100 != 0) or (ano_selecionado % 400 == 0) else 28

dia_fim = st.sidebar.number_input(
    "Dia (até qual dia do mês)",
    min_value=1,
    max_value=dia_max,
    value=min(hoje.day, dia_max) if mes_selecionado == hoje.month else dia_max
)

# Conexão com banco
conn = get_db_connection()

if conn:
    # Buscar supervisores
    df_supervisores = buscar_supervisores(conn)
    
    if not df_supervisores.empty:
        supervisores_opcoes = df_supervisores['codsupervisor'].tolist()
        supervisores_nomes = {row['codsupervisor']: row['nome'] for _, row in df_supervisores.iterrows()}
        
        # Seleção de supervisores (1 e 2 padrão)
        supervisores_selecionados = st.sidebar.multiselect(
            "Supervisores",
            options=supervisores_opcoes,
            default=[1, 2] if 1 in supervisores_opcoes and 2 in supervisores_opcoes else supervisores_opcoes[:2] if len(supervisores_opcoes) >= 2 else supervisores_opcoes,
            format_func=lambda x: f"{x} - {supervisores_nomes.get(x, '')}"
        )
    else:
        supervisores_selecionados = [1, 2]
    
    # Botão para atualizar dados
    if st.sidebar.button("🔄 Atualizar Dados", type="primary"):
        st.cache_resource.clear()
        st.rerun()
    
    # Calcular mês anterior para exibição
    if mes_selecionado == 1:
        mes_anterior_display = 12
        ano_anterior_display = ano_selecionado - 1
    else:
        mes_anterior_display = mes_selecionado - 1
        ano_anterior_display = ano_selecionado
    
    # Dia máximo do mês anterior
    dia_max_anterior_display = 31
    if mes_anterior_display in [4, 6, 9, 11]:
        dia_max_anterior_display = 30
    elif mes_anterior_display == 2:
        dia_max_anterior_display = 29 if (ano_anterior_display % 4 == 0 and ano_anterior_display % 100 != 0) or (ano_anterior_display % 400 == 0) else 28
    dia_fim_anterior_display = min(dia_fim, dia_max_anterior_display)
    
    # Buscar dados do mês atual
    st.sidebar.info(f"📅 Período Atual: 01/{mes_selecionado:02d}/{ano_selecionado} até {dia_fim:02d}/{mes_selecionado:02d}/{ano_selecionado}")
    st.sidebar.info(f"📅 Período Anterior: 01/{mes_anterior_display:02d}/{ano_anterior_display} até {dia_fim_anterior_display:02d}/{mes_anterior_display:02d}/{ano_anterior_display}")
    
    with st.spinner("Buscando dados do mês atual..."):
        df_mes_atual = buscar_dados_vendedores(conn, mes_selecionado, ano_selecionado, dia_fim, supervisores_selecionados)
    
    # Calcular mês anterior (usar o mesmo dia para comparação justa)
    if mes_selecionado == 1:
        mes_anterior = 12
        ano_anterior = ano_selecionado - 1
    else:
        mes_anterior = mes_selecionado - 1
        ano_anterior = ano_selecionado
    
    # Dia máximo do mês anterior (para garantir que não ultrapasse)
    dia_max_anterior = 31
    if mes_anterior in [4, 6, 9, 11]:
        dia_max_anterior = 30
    elif mes_anterior == 2:
        dia_max_anterior = 29 if (ano_anterior % 4 == 0 and ano_anterior % 100 != 0) or (ano_anterior % 400 == 0) else 28
    
    # Usar o mesmo dia do mês anterior (comparação justa)
    # Se dia_fim é 05, comparar até dia 05 do mês anterior também
    dia_fim_anterior = min(dia_fim, dia_max_anterior)
    
    with st.spinner(f"Buscando dados do mês anterior (até dia {dia_fim_anterior})..."):
        df_mes_anterior = buscar_dados_vendedores(conn, mes_anterior, ano_anterior, dia_fim_anterior, supervisores_selecionados)
    
    if not df_mes_atual.empty:
        # Renomear colunas do mês atual
        df_mes_atual_renamed = df_mes_atual.rename(columns={
            'venda_mes': 'venda_mes_atual',
            'venda_pedido': 'venda_pedido_atual',
            'margem_percent': 'margem_percent_atual',
            'positivacao': 'positivacao_atual',
            'mix': 'mix_atual',
            'qtd_nf': 'qtd_nf_atual'
        })
        
        # Renomear colunas do mês anterior
        if not df_mes_anterior.empty:
            df_mes_anterior_renamed = df_mes_anterior.rename(columns={
                'venda_mes': 'venda_mes_anterior',
                'margem_percent': 'margem_percent_anterior',
                'positivacao': 'positivacao_anterior',
                'mix': 'mix_anterior',
                'qtd_nf': 'qtd_nf_anterior'
            })
            
            # Combinar dados por RCA e CODCLI
            df_combinado = df_mes_atual_renamed.merge(
                df_mes_anterior_renamed[['rca', 'codcli', 'venda_mes_anterior', 'margem_percent_anterior', 'positivacao_anterior', 'mix_anterior', 'qtd_nf_anterior']],
                on=['rca', 'codcli'],
                how='left'
            )
        else:
            df_combinado = df_mes_atual_renamed.copy()
            df_combinado['venda_mes_anterior'] = 0
            df_combinado['margem_percent_anterior'] = 0
            df_combinado['positivacao_anterior'] = 0
            df_combinado['mix_anterior'] = 0
            df_combinado['qtd_nf_anterior'] = 0
        
        # Preencher NaN com 0
        df_combinado = df_combinado.fillna(0)
        
        # Criar DataFrame final
        df_final = pd.DataFrame({
            'RCA': df_combinado['rca'].astype(int),
            'Vendedor': df_combinado['nome_vendedor'].fillna(''),
            'Cod Cliente': df_combinado['codcli'].astype(int),
            'Cliente': df_combinado['nome_cliente'].fillna(''),
            'Venda Mês Anterior': df_combinado['venda_mes_anterior'].round(2),
            'Venda Mês Atual': df_combinado['venda_mes_atual'].round(2),
            'Venda + Pedido (F+L)': df_combinado['venda_pedido_atual'].round(2),
            'Margem % Atual': df_combinado['margem_percent_atual'].round(2),
            'Margem % Mês Anterior': df_combinado['margem_percent_anterior'].round(2),
            'Positivação': df_combinado['positivacao_atual'].astype(int),
            'Positivação Mês Anterior': df_combinado['positivacao_anterior'].astype(int),
            'Mix Atual': df_combinado['mix_atual'].astype(int),
            'Mix Anterior': df_combinado['mix_anterior'].astype(int),
            'Quant NF Atual': df_combinado['qtd_nf_atual'].astype(int),
            'Quant NF Anterior': df_combinado['qtd_nf_anterior'].astype(int)
        })
        
        # Ordenar por venda mês atual
        df_final = df_final.sort_values('Venda Mês Atual', ascending=False).reset_index(drop=True)
        
        # =================================================================
        # CARDS DE TOTAIS COM INDICADORES
        # =================================================================
        st.header("📈 Totais e Indicadores")
        
        # Calcular totais
        total_venda_atual = df_final['Venda Mês Atual'].sum()
        total_venda_anterior = df_final['Venda Mês Anterior'].sum()
        total_venda_pedido = df_final['Venda + Pedido (F+L)'].sum()
        
        # Calcular margem média ponderada
        margem_media_atual = ((df_final['Venda Mês Atual'] * df_final['Margem % Atual']).sum() / total_venda_atual) if total_venda_atual > 0 else 0
        margem_media_anterior = ((df_final['Venda Mês Anterior'] * df_final['Margem % Mês Anterior']).sum() / total_venda_anterior) if total_venda_anterior > 0 else 0
        
        total_positivacao_atual = df_final['Positivação'].sum()
        total_positivacao_anterior = df_final['Positivação Mês Anterior'].sum()
        
        total_mix_atual = df_final['Mix Atual'].sum()
        total_mix_anterior = df_final['Mix Anterior'].sum()
        
        total_nf_atual = df_final['Quant NF Atual'].sum()
        total_nf_anterior = df_final['Quant NF Anterior'].sum()
        
        # Função para calcular variação e ícone
        def calcular_variacao(atual, anterior):
            if anterior == 0:
                return 100.0 if atual > 0 else 0.0, "🟢" if atual > 0 else "➡️"
            variacao = ((atual - anterior) / anterior) * 100
            if variacao > 0:
                return variacao, "🟢"
            elif variacao < 0:
                return variacao, "🔴"
            else:
                return 0.0, "➡️"
        
        # Cards em colunas
        col1, col2, col3, col4 = st.columns(4)
        
        with col1:
            var_venda, icon_venda = calcular_variacao(total_venda_atual, total_venda_anterior)
            st.metric(
                "💰 Faturamento Mês Atual",
                f"R$ {total_venda_atual:,.2f}".replace(",", ".").replace(".", ",", 1),
                f"{icon_venda} {var_venda:+.2f}%"
            )
        
        with col2:
            var_margem, icon_margem = calcular_variacao(margem_media_atual, margem_media_anterior)
            st.metric(
                "📊 Margem % Média",
                f"{margem_media_atual:.2f}%",
                f"{icon_margem} {var_margem:+.2f}%"
            )
        
        with col3:
            var_positivacao, icon_positivacao = calcular_variacao(total_positivacao_atual, total_positivacao_anterior)
            st.metric(
                "✅ Positivação",
                f"{total_positivacao_atual:,.0f}".replace(",", "."),
                f"{icon_positivacao} {var_positivacao:+.2f}%"
            )
        
        with col4:
            var_mix, icon_mix = calcular_variacao(total_mix_atual, total_mix_anterior)
            st.metric(
                "📦 Mix Total",
                f"{total_mix_atual:,.0f}".replace(",", "."),
                f"{icon_mix} {var_mix:+.2f}%"
            )
        
        # Segunda linha de cards
        col5, col6, col7, col8 = st.columns(4)
        
        with col5:
            st.metric(
                "💼 Venda + Pedido (F+L)",
                f"R$ {total_venda_pedido:,.2f}".replace(",", ".").replace(".", ",", 1)
            )
        
        with col6:
            var_nf, icon_nf = calcular_variacao(total_nf_atual, total_nf_anterior)
            st.metric(
                "📄 Quantidade de NFs",
                f"{total_nf_atual:,.0f}".replace(",", "."),
                f"{icon_nf} {var_nf:+.2f}%"
            )
        
        with col7:
            st.metric(
                "📅 Período",
                f"{dia_fim:02d}/{mes_selecionado:02d}/{ano_selecionado}"
            )
        
        with col8:
            st.metric(
                "👥 Supervisores",
                f"{len(supervisores_selecionados)} selecionados"
            )
        
        # =================================================================
        # TABELA
        # =================================================================
        st.header("📋 Tabela Detalhada")
        
        # Formatação da tabela para exibição
        df_display = df_final.copy()
        df_display['Venda Mês Anterior'] = df_display['Venda Mês Anterior'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        df_display['Venda Mês Atual'] = df_display['Venda Mês Atual'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        df_display['Venda + Pedido (F+L)'] = df_display['Venda + Pedido (F+L)'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        df_display['Margem % Atual'] = df_display['Margem % Atual'].apply(lambda x: f"{x:.2f}%")
        df_display['Margem % Mês Anterior'] = df_display['Margem % Mês Anterior'].apply(lambda x: f"{x:.2f}%")
        
        st.dataframe(df_display, use_container_width=True, height=600)
        
        # Informação sobre quantidade de registros
        st.info(f"📊 Total de registros: {len(df_final):,} linhas (uma por vendedor/cliente)")
        
        # =================================================================
        # BOTÃO PARA BAIXAR EXCEL
        # =================================================================
        st.header("💾 Exportar Dados")
        
        def gerar_excel(df):
            """Gera Excel formatado"""
            wb = Workbook()
            ws = wb.active
            ws.title = "Análise Vendedores"
            
            # Estilos
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin'),
            )
            
            # Cabeçalho
            headers = list(df.columns)
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            # Dados
            for row_idx, row_data in enumerate(df.values, 2):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = value
                    cell.border = border
                    
                    # Formatação de números
                    if 'Venda' in headers[col_idx - 1]:
                        cell.number_format = 'R$ #,##0.00'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    elif 'Margem %' in headers[col_idx - 1]:
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    elif 'Positivação' in headers[col_idx - 1] or 'Mix' in headers[col_idx - 1] or 'Quant NF' in headers[col_idx - 1]:
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    elif 'RCA' in headers[col_idx - 1] or 'Cod Cliente' in headers[col_idx - 1]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')
            
            # Ajustar larguras
            for col_idx, header in enumerate(headers, 1):
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = max(len(str(header)) + 2, 15)
            
            # Congelar primeira linha
            ws.freeze_panes = 'A2'
            
            # Salvar em buffer
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer
        
        excel_buffer = gerar_excel(df_final)
        
        st.download_button(
            label="📥 Baixar Excel",
            data=excel_buffer,
            file_name=f"analise_vendedores_{mes_selecionado:02d}_{ano_selecionado}_{dia_fim:02d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
    else:
        st.warning("Nenhum dado encontrado para os filtros selecionados.")
else:
    st.error("Não foi possível conectar ao banco de dados. Verifique as configurações.")

