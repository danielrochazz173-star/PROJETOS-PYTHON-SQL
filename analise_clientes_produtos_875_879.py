"""
Script de análise de clientes que compraram produtos 875 ou 879 a R$ 9,90
Permite escolher o mês/ano para análise
"""

import oracledb
import pandas as pd
from datetime import datetime
from calendar import monthrange
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configuração do Oracle Instant Client
oracledb.init_oracle_client(lib_dir=r"C:\oracle\instantclient_23_9")

# =========================================================================
# CONFIGURAÇÕES
# =========================================================================
DB_HOST = '10.0.0.10'
DB_PORT = 1521
DB_SERVICE = 'PROD'
DB_USER = 'powerbi'
DB_PASSWORD = 'cbjc4xp3nlq6'

# Produtos a buscar
CODPROD_LIST = [875, 879]
PRECO_BUSCA = 9.90

# =========================================================================
# FUNÇÕES
# =========================================================================
def get_db_connection():
    """Conecta ao banco Oracle"""
    try:
        return oracledb.connect(user=DB_USER, password=DB_PASSWORD, dsn=f"{DB_HOST}:{DB_PORT}/{DB_SERVICE}")
    except oracledb.DatabaseError as e:
        print(f"❌ Erro ao conectar ao Oracle: {e}")
        return None

def get_query_clientes(mes, ano):
    """Gera a query SQL para buscar clientes"""
    data_inicio = f"01/{mes:02d}/{ano}"
    # Calcula o último dia do mês
    ultimo_dia = monthrange(ano, mes)[1]
    data_fim = f"{ultimo_dia:02d}/{mes:02d}/{ano}"
    
    codprod_str = ', '.join(map(str, CODPROD_LIST))
    
    query = f"""
    SELECT 
        CL.CODCLI,
        CL.CLIENTE,
        C.NUMNOTA,
        C.NUMPED,
        C.DATA,
        I.CODPROD,
        P.DESCRICAO AS PRODUTO,
        I.PVENDA,
        I.QT,
        ROUND(I.QT * I.PVENDA, 2) AS VALOR_TOTAL,
        C.CODUSUR,
        U.NOME AS VENDEDOR
    FROM PCPEDC C
    JOIN PCPEDI I ON C.NUMPED = I.NUMPED
    JOIN PCCLIENT CL ON C.CODCLI = CL.CODCLI
    JOIN PCPRODUT P ON I.CODPROD = P.CODPROD
    LEFT JOIN PCUSUARI U ON C.CODUSUR = U.CODUSUR
    WHERE C.DATA BETWEEN TO_DATE('{data_inicio}', 'DD/MM/YYYY') 
        AND TO_DATE('{data_fim}', 'DD/MM/YYYY')
        AND I.CODPROD IN ({codprod_str})
        AND ROUND(I.PVENDA, 2) = {PRECO_BUSCA}
        AND C.CODFILIAL IN ('1', '98')
        AND C.POSICAO = 'F'
        AND C.DTCANCEL IS NULL
        AND C.CONDVENDA NOT IN (4, 8, 10, 13, 20, 98, 99)
        AND NVL(I.BONIFIC, 'N') = 'N'
    ORDER BY C.DATA DESC, CL.CODCLI, I.CODPROD
    """
    
    return query

# =========================================================================
# EXECUÇÃO
# =========================================================================
if __name__ == "__main__":
    print("=" * 100)
    print("📊 ANÁLISE DE CLIENTES - PRODUTOS 875 E 879 A R$ 9,90")
    print("=" * 100)
    print()
    
    # Solicita mês e ano
    try:
        ano = int(input("Digite o ano (ex: 2025): "))
        mes = int(input("Digite o mês (1-12): "))
        
        if mes < 1 or mes > 12:
            print("❌ Mês inválido. Deve ser entre 1 e 12.")
            exit(1)
        
        print()
        print(f"🔍 Buscando clientes que compraram produtos {CODPROD_LIST} a R$ {PRECO_BUSCA:.2f}")
        print(f"   Período: {mes:02d}/{ano}")
        print()
        
    except ValueError:
        print("❌ Entrada inválida. Use apenas números.")
        exit(1) 
    
    # Conecta ao banco
    print("🔌 Conectando ao Oracle...")
    conn = get_db_connection()
    
    if not conn:
        print("❌ Falha na conexão. Abortando.")
        exit(1)
    
    try:
        # Executa a query
        print("📈 Executando query...")
        query = get_query_clientes(mes, ano)
        df = pd.read_sql_query(query, conn)
        
        if df.empty:
            print("⚠️  Nenhum cliente encontrado com os critérios especificados.")
            exit(0)
        
        print(f"✅ {len(df)} registro(s) encontrado(s).\n")
        print("=" * 120)
        print(f"{'CÓD.':<8} {'CLIENTE':<40} {'NF':<10} {'PEDIDO':<10} {'DATA':<12} {'PROD.':<6} {'PRODUTO':<25} {'QTD':<8} {'VALOR':<12}")
        print("=" * 120)
        
        # Mostra todos os registros
        for idx, row in df.iterrows():
            codcli = str(row['CODCLI'])[:8]
            cliente = str(row['CLIENTE'])[:38]
            numnota = str(row['NUMNOTA'])[:10] if pd.notna(row['NUMNOTA']) else 'N/A'
            numped = str(row['NUMPED'])[:10]
            data = row['DATA'].strftime('%d/%m/%Y') if pd.notna(row['DATA']) else 'N/A'
            codprod = str(row['CODPROD'])[:6]
            produto = str(row['PRODUTO'])[:23]
            qt = f"{row['QT']:.0f}" if pd.notna(row['QT']) else '0'
            valor = f"R$ {row['VALOR_TOTAL']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            
            print(f"{codcli:<8} {cliente:<40} {numnota:<10} {numped:<10} {data:<12} {codprod:<6} {produto:<25} {qt:<8} {valor:<12}")
        
        print("=" * 120)
        print(f"\n📊 RESUMO:")
        print(f"   Total de registros: {len(df)}")
        print(f"   Clientes únicos: {df['CODCLI'].nunique()}")
        print(f"   Pedidos únicos: {df['NUMPED'].nunique()}")
        print(f"   Valor total: R$ {df['VALOR_TOTAL'].sum():,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
        print()
        
        # Agrupa por cliente
        print("=" * 120)
        print("📋 RESUMO POR CLIENTE:")
        print("=" * 120)
        print(f"{'CÓD.':<8} {'CLIENTE':<50} {'QTD PEDIDOS':<15} {'VALOR TOTAL':<15}")
        print("=" * 100)
        
        resumo_cliente = df.groupby(['CODCLI', 'CLIENTE']).agg({
            'NUMPED': 'nunique',
            'VALOR_TOTAL': 'sum'
        }).reset_index()
        resumo_cliente = resumo_cliente.sort_values('VALOR_TOTAL', ascending=False)
        
        for idx, row in resumo_cliente.iterrows():
            codcli = str(row['CODCLI'])[:8]
            cliente = str(row['CLIENTE'])[:48]
            qtd_pedidos = str(row['NUMPED'])
            valor_total = f"R$ {row['VALOR_TOTAL']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            
            print(f"{codcli:<8} {cliente:<50} {qtd_pedidos:<15} {valor_total:<15}")
        
        print("=" * 100)
        print()
        
        # Salva em Excel formatado
        arquivo_excel = f"clientes_produtos_875_879_{mes:02d}_{ano}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        # Renomeia colunas para o Excel
        df_excel = df.rename(columns={'NUMNOTA': 'NF'})
        
        with pd.ExcelWriter(arquivo_excel, engine='openpyxl') as writer:
            # Aba 1: Detalhado
            df_excel.to_excel(writer, sheet_name='Detalhado', index=False)
            ws_detalhado = writer.sheets['Detalhado']
            
            # Formata cabeçalho
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for cell in ws_detalhado[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            # Ajusta largura das colunas
            ws_detalhado.column_dimensions['A'].width = 10  # CODCLI
            ws_detalhado.column_dimensions['B'].width = 40  # CLIENTE
            ws_detalhado.column_dimensions['C'].width = 12  # NF
            ws_detalhado.column_dimensions['D'].width = 12  # NUMPED
            ws_detalhado.column_dimensions['E'].width = 12  # DATA
            ws_detalhado.column_dimensions['F'].width = 10  # CODPROD
            ws_detalhado.column_dimensions['G'].width = 35  # PRODUTO
            ws_detalhado.column_dimensions['H'].width = 10  # PVENDA
            ws_detalhado.column_dimensions['I'].width = 10  # QT
            ws_detalhado.column_dimensions['J'].width = 15  # VALOR_TOTAL
            ws_detalhado.column_dimensions['K'].width = 10  # CODUSUR
            ws_detalhado.column_dimensions['L'].width = 25  # VENDEDOR
            
            # Formata dados
            for row in ws_detalhado.iter_rows(min_row=2, max_row=ws_detalhado.max_row):
                for cell in row:
                    cell.border = border
                    if cell.column == 10:  # VALOR_TOTAL
                        cell.number_format = 'R$ #,##0.00'
                    elif cell.column == 8:  # PVENDA
                        cell.number_format = 'R$ #,##0.00'
                    elif cell.column == 5:  # DATA
                        cell.number_format = 'DD/MM/YYYY'
            
            # Aba 2: Resumo por Cliente
            resumo_cliente.to_excel(writer, sheet_name='Resumo por Cliente', index=False)
            ws_resumo = writer.sheets['Resumo por Cliente']
            
            # Formata cabeçalho do resumo
            for cell in ws_resumo[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            # Ajusta largura das colunas do resumo
            ws_resumo.column_dimensions['A'].width = 10  # CODCLI
            ws_resumo.column_dimensions['B'].width = 40  # CLIENTE
            ws_resumo.column_dimensions['C'].width = 15  # NUMPED (qtd pedidos)
            ws_resumo.column_dimensions['D'].width = 15  # VALOR_TOTAL
            
            # Formata dados do resumo
            for row in ws_resumo.iter_rows(min_row=2, max_row=ws_resumo.max_row):
                for cell in row:
                    cell.border = border
                    if cell.column == 4:  # VALOR_TOTAL
                        cell.number_format = 'R$ #,##0.00'
            
            # Congela primeira linha
            ws_detalhado.freeze_panes = 'A2'
            ws_resumo.freeze_panes = 'A2'
        
        print(f"💾 Dados salvos em Excel: {arquivo_excel}")
        
    except Exception as e:
        print(f"❌ Erro ao executar query: {e}")
        import traceback
        traceback.print_exc()
    finally:
        conn.close()
        print("\n🔌 Conexão fechada.")

