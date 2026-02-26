"""
Script de análise dos últimos 3 meses (outubro, novembro, dezembro)
para os produtos 1615 (gra file) e 543 (gra file picanha)
Mostra valor vendido e margem por mês, exportando para Excel
"""

import oracledb
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configuração do Oracle Instant Client
oracledb.init_oracle_client(lib_dir=r"C:\oracle\instantclient_23_9")

# =========================================================================
# CONFIGURAÇÕES
# =========================================================================
DB_HOST = '10.0.0.10'
DB_PORT = 1521
DB_SERVICE = 'PROD'
DB_USER = 'powerbi'
DB_PASSWORD = 'cbjc4xp3nlq6'

# Produtos a analisar
CODPROD_LIST = [1615, 543]

# Meses a analisar (outubro, novembro, dezembro)
MESES_ANALISE = [
    {'mes': 10, 'ano': 2025, 'nome': 'Outubro'},
    {'mes': 11, 'ano': 2025, 'nome': 'Novembro'},
    {'mes': 12, 'ano': 2025, 'nome': 'Dezembro'}
]

# =========================================================================
# FUNÇÕES
# =========================================================================
def get_db_connection():
    """Conecta ao banco Oracle"""
    try:
        return oracledb.connect(user=DB_USER, password=DB_PASSWORD, dsn=f"{DB_HOST}:{DB_PORT}/{DB_SERVICE}")
    except oracledb.DatabaseError as e:
        print(f"❌ Erro ao conectar ao Oracle: {e}")
        return None

def buscar_vendas_mes(conn, mes, ano, codprod_list):
    """Busca vendas de um mês específico para os produtos informados
    Faturamento líquido (faturamento bruto - devoluções)
    Filtra apenas supervisores 1 e 2
    """
    from calendar import monthrange
    
    data_inicio = f"01/{mes:02d}/{ano}"
    ultimo_dia = monthrange(ano, mes)[1]
    data_fim = f"{ultimo_dia:02d}/{mes:02d}/{ano}"
    
    codprod_str = ', '.join(map(str, codprod_list))
    
    query = f"""
    WITH VENDAS_VALIDAS AS (
        SELECT 
            I.CODPROD,
            C.CODFILIAL,
            SUM(
                CASE 
                    WHEN NVL(I.BONIFIC, 'N') = 'N' THEN 
                        DECODE(C.CONDVENDA, 
                            5, 0, 6, 0, 11, 0, 12, 0, 
                            ROUND(NVL(I.QT, 0) * NVL(I.PVENDA, 0), 2)
                        )
                    ELSE 0 
                END
            ) AS VALOR_BRUTO,
            SUM(
                CASE 
                    WHEN NVL(I.BONIFIC, 'N') = 'N' THEN 
                        DECODE(C.CONDVENDA, 
                            5, 0, 6, 0, 11, 0, 12, 0, 
                            NVL(I.QT, 0) * NVL(I.VLCUSTOFIN, 0)
                        )
                    ELSE 0 
                END
            ) AS CUSTO_BRUTO
        FROM PCPEDC C 
        JOIN PCPEDI I ON C.NUMPED = I.NUMPED
        JOIN PCUSUARI U ON C.CODUSUR = U.CODUSUR
        WHERE C.DATA BETWEEN TO_DATE('{data_inicio}', 'DD/MM/YYYY') 
            AND TO_DATE('{data_fim}', 'DD/MM/YYYY')
            AND I.CODPROD IN ({codprod_str})
            AND C.CODFILIAL IN ('1', '98')
            AND C.POSICAO = 'F'
            AND C.DTCANCEL IS NULL
            AND C.CONDVENDA NOT IN (4, 8, 10, 13, 20, 98, 99)
            AND U.CODSUPERVISOR IN (1, 2)
        GROUP BY I.CODPROD, C.CODFILIAL
    ),
    DEVOLUCOES AS (
        SELECT 
            D.CODPROD,
            D.CODFILIAL,
            SUM(NVL(D.VLDEVOLUCAO, 0)) AS VALOR_DEVOLVIDO,
            SUM(NVL(D.VLCUSTOFIN, 0)) AS CUSTO_DEVOLVIDO
        FROM VIEW_DEVOL_RESUMO_FATURAMENTO D
        JOIN PCUSUARI U ON D.CODUSUR = U.CODUSUR
        WHERE D.DTENT BETWEEN TO_DATE('{data_inicio}', 'DD/MM/YYYY') 
            AND TO_DATE('{data_fim}', 'DD/MM/YYYY')
            AND D.CODPROD IN ({codprod_str})
            AND D.CODFILIAL IN ('1', '98')
            AND U.CODSUPERVISOR IN (1, 2)
        GROUP BY D.CODPROD, D.CODFILIAL
    )
    SELECT 
        COALESCE(V.CODPROD, D.CODPROD) AS CODPROD,
        COALESCE(V.CODFILIAL, D.CODFILIAL) AS CODFILIAL,
        (NVL(V.VALOR_BRUTO, 0) - NVL(D.VALOR_DEVOLVIDO, 0)) AS VALOR_VENDIDO,
        (NVL(V.CUSTO_BRUTO, 0) - NVL(D.CUSTO_DEVOLVIDO, 0)) AS CUSTO_TOTAL,
        (NVL(V.VALOR_BRUTO, 0) - NVL(D.VALOR_DEVOLVIDO, 0) - 
         NVL(V.CUSTO_BRUTO, 0) + NVL(D.CUSTO_DEVOLVIDO, 0)) AS MARGEM_VALOR
    FROM VENDAS_VALIDAS V
    FULL OUTER JOIN DEVOLUCOES D ON V.CODPROD = D.CODPROD AND V.CODFILIAL = D.CODFILIAL
    WHERE (NVL(V.VALOR_BRUTO, 0) > 0 OR NVL(D.VALOR_DEVOLVIDO, 0) > 0)
    """
    
    df = pd.read_sql_query(query, conn)
    df.columns = [x.lower() for x in df.columns]
    
    # Calcula margem percentual
    if not df.empty:
        df['margem_percentual'] = (df['margem_valor'] / df['valor_vendido'] * 100).fillna(0)
        # Garante que produtos sem vendas tenham valores zerados
        df['valor_vendido'] = df['valor_vendido'].fillna(0)
        df['custo_total'] = df['custo_total'].fillna(0)
        df['margem_valor'] = df['margem_valor'].fillna(0)
    
    return df

def buscar_descricao_produtos(conn, codprod_list):
    """Busca descrição dos produtos"""
    codprod_str = ', '.join(map(str, codprod_list))
    
    query = f"""
    SELECT CODPROD, DESCRICAO
    FROM PCPRODUT
    WHERE CODPROD IN ({codprod_str})
    """
    
    df = pd.read_sql_query(query, conn)
    df.columns = [x.lower() for x in df.columns]
    return df


# =========================================================================
# EXECUÇÃO
# =========================================================================
if __name__ == "__main__":
    print("=" * 100)
    print("📊 ANÁLISE DE PRODUTOS 1615 E 543 - ÚLTIMOS 3 MESES")
    print("   Filtro: Apenas Supervisores 1 e 2")
    print("   Faturamento: Líquido (faturamento bruto - devoluções)")
    print("=" * 100)
    print()
    
    # Conecta ao banco
    print("🔌 Conectando ao Oracle...")
    conn = get_db_connection()
    
    if not conn:
        print("❌ Falha na conexão. Abortando.")
        exit(1)
    
    try:
        # Busca descrições dos produtos
        print("📦 Buscando descrições dos produtos...")
        df_produtos = buscar_descricao_produtos(conn, CODPROD_LIST)
        
        if df_produtos.empty:
            print("⚠️  Nenhum produto encontrado.")
            exit(0)
        
        # Cria dicionário com descrições
        descricoes = {}
        for _, row in df_produtos.iterrows():
            descricoes[int(row['codprod'])] = row['descricao']
        
        print(f"✅ {len(df_produtos)} produto(s) encontrado(s).\n")
        
        # Busca vendas de cada mês
        dados_meses = {}
        
        for mes_info in MESES_ANALISE:
            mes = mes_info['mes']
            ano = mes_info['ano']
            nome_mes = mes_info['nome']
            
            print(f"📈 Buscando vendas de {nome_mes}/{ano}...")
            df_vendas = buscar_vendas_mes(conn, mes, ano, CODPROD_LIST)
            
            # Processa dados por produto e filial
            for _, row in df_vendas.iterrows():
                codprod = int(row['codprod'])
                codfilial = str(row['codfilial']) if pd.notna(row['codfilial']) else '1'
                valor_vendido = float(row['valor_vendido']) if pd.notna(row['valor_vendido']) else 0.0
                custo_total = float(row['custo_total']) if pd.notna(row['custo_total']) else 0.0
                margem_valor = float(row['margem_valor']) if pd.notna(row['margem_valor']) else 0.0
                margem_percentual = float(row['margem_percentual']) if pd.notna(row['margem_percentual']) else 0.0
                
                if codprod not in dados_meses:
                    dados_meses[codprod] = {}
                
                if codfilial not in dados_meses[codprod]:
                    dados_meses[codprod][codfilial] = {}
                
                dados_meses[codprod][codfilial][nome_mes] = {
                    'valor_vendido': valor_vendido,
                    'margem_valor': margem_valor,
                    'margem_percentual': margem_percentual
                }
            
            print(f"   ✅ {len(df_vendas)} registro(s) encontrado(s).")
        
        print()
        
        # Prepara dados para Excel
        print("📊 Preparando dados para Excel...")
        
        # Cria estrutura de dados para o Excel (uma linha por produto/filial)
        dados_excel = []
        
        for codprod in CODPROD_LIST:
            descricao = descricoes.get(codprod, f"Produto {codprod}")
            
            # Para cada filial (1 e 98)
            for filial in ['1', '98']:
                # Busca dados de cada mês para esta filial
                outubro = dados_meses.get(codprod, {}).get(filial, {}).get('Outubro', {'valor_vendido': 0, 'margem_valor': 0, 'margem_percentual': 0})
                novembro = dados_meses.get(codprod, {}).get(filial, {}).get('Novembro', {'valor_vendido': 0, 'margem_valor': 0, 'margem_percentual': 0})
                dezembro = dados_meses.get(codprod, {}).get(filial, {}).get('Dezembro', {'valor_vendido': 0, 'margem_valor': 0, 'margem_percentual': 0})
                
                dados_excel.append({
                    'CODPROD': codprod,
                    'PRODUTO': descricao,
                    'FILIAL': filial,
                    'OUTUBRO_VALOR': outubro['valor_vendido'],
                    'OUTUBRO_MARGEM': outubro['margem_percentual'],
                    'NOVEMBRO_VALOR': novembro['valor_vendido'],
                    'NOVEMBRO_MARGEM': novembro['margem_percentual'],
                    'DEZEMBRO_VALOR': dezembro['valor_vendido'],
                    'DEZEMBRO_MARGEM': dezembro['margem_percentual']
                })
        
        # Cria DataFrame
        df_excel = pd.DataFrame(dados_excel)
        
        # Calcula análises para as abas extras
        # Análise de crescimento
        analise_crescimento = []
        for _, row in df_excel.iterrows():
            out_valor = row['OUTUBRO_VALOR']
            nov_valor = row['NOVEMBRO_VALOR']
            dez_valor = row['DEZEMBRO_VALOR']
            
            cres_out_nov = ((nov_valor - out_valor) / out_valor * 100) if out_valor > 0 else 0
            cres_nov_dez = ((dez_valor - nov_valor) / nov_valor * 100) if nov_valor > 0 else 0
            cres_total = ((dez_valor - out_valor) / out_valor * 100) if out_valor > 0 else 0
            
            analise_crescimento.append({
                'Código': row['CODPROD'],
                'Produto': row['PRODUTO'],
                'Filial': row['FILIAL'],
                'Outubro': out_valor,
                'Novembro': nov_valor,
                'Dezembro': dez_valor,
                'Cresc. Out→Nov %': cres_out_nov,
                'Cresc. Nov→Dez %': cres_nov_dez,
                'Cresc. Total %': cres_total
            })
        
        df_crescimento = pd.DataFrame(analise_crescimento)
        df_crescimento = df_crescimento.sort_values('Cresc. Total %', ascending=False)
        
        # Análise por filial
        analise_filial = []
        for filial in ['1', '98']:
            df_filial = df_excel[df_excel['FILIAL'] == filial]
            out_total = df_filial['OUTUBRO_VALOR'].sum()
            nov_total = df_filial['NOVEMBRO_VALOR'].sum()
            dez_total = df_filial['DEZEMBRO_VALOR'].sum()
            
            cres_out_nov = ((nov_total - out_total) / out_total * 100) if out_total > 0 else 0
            cres_nov_dez = ((dez_total - nov_total) / nov_total * 100) if nov_total > 0 else 0
            cres_total = ((dez_total - out_total) / out_total * 100) if out_total > 0 else 0
            
            analise_filial.append({
                'Filial': filial,
                'Outubro': out_total,
                'Novembro': nov_total,
                'Dezembro': dez_total,
                'Cresc. Out→Nov %': cres_out_nov,
                'Cresc. Nov→Dez %': cres_nov_dez,
                'Cresc. Total %': cres_total
            })
        
        df_filial = pd.DataFrame(analise_filial)
        
        # Salva em Excel formatado
        arquivo_excel = f"analise_produtos_1615_543_3meses_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        with pd.ExcelWriter(arquivo_excel, engine='openpyxl') as writer:
            # Aba 1: Dados Detalhados - Organizado por Filial
            ws_detalhado = writer.book.create_sheet('Dados Detalhados')
            
            # Cabeçalho
            header = ['Código', 'Produto', 'Outubro - Valor', 'Outubro - Margem %', 
                     'Novembro - Valor', 'Novembro - Margem %', 
                     'Dezembro - Valor', 'Dezembro - Margem %']
            
            row_num = 1
            
            # Organiza por filial
            for filial in ['1', '98']:
                # Cabeçalho da filial
                ws_detalhado.merge_cells(f'A{row_num}:H{row_num}')
                cell_filial = ws_detalhado.cell(row=row_num, column=1)
                cell_filial.value = f'FILIAL {filial}'
                cell_filial.fill = PatternFill(start_color="D40000", end_color="D40000", fill_type="solid")
                cell_filial.font = Font(bold=True, color="FFFFFF", size=14)
                cell_filial.alignment = Alignment(horizontal='center', vertical='center')
                row_num += 1
                
                # Cabeçalho das colunas
                for col_num, header_text in enumerate(header, 1):
                    cell = ws_detalhado.cell(row=row_num, column=col_num)
                    cell.value = header_text
                    cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF", size=11)
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )
                row_num += 1
                
                # Dados da filial
                df_filial = df_excel[df_excel['FILIAL'] == filial].copy()
                df_filial = df_filial.sort_values('CODPROD')
                
                for _, row in df_filial.iterrows():
                    ws_detalhado.cell(row=row_num, column=1, value=int(row['CODPROD']))
                    ws_detalhado.cell(row=row_num, column=2, value=row['PRODUTO'])
                    ws_detalhado.cell(row=row_num, column=3, value=row['OUTUBRO_VALOR'])
                    ws_detalhado.cell(row=row_num, column=4, value=row['OUTUBRO_MARGEM'])
                    ws_detalhado.cell(row=row_num, column=5, value=row['NOVEMBRO_VALOR'])
                    ws_detalhado.cell(row=row_num, column=6, value=row['NOVEMBRO_MARGEM'])
                    ws_detalhado.cell(row=row_num, column=7, value=row['DEZEMBRO_VALOR'])
                    ws_detalhado.cell(row=row_num, column=8, value=row['DEZEMBRO_MARGEM'])
                    
                    # Formata células
                    for col in range(1, 9):
                        cell = ws_detalhado.cell(row=row_num, column=col)
                        cell.border = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )
                        if col == 1:  # Código
                            cell.alignment = Alignment(horizontal='center', vertical='center')
                        elif col == 2:  # Produto
                            cell.alignment = Alignment(horizontal='left', vertical='center')
                        elif col in [3, 5, 7]:  # Valores
                            cell.number_format = 'R$ #,##0.00'
                            cell.alignment = Alignment(horizontal='right', vertical='center')
                        elif col in [4, 6, 8]:  # Margens
                            cell.number_format = '0.00"%"'
                            cell.alignment = Alignment(horizontal='right', vertical='center')
                    
                    row_num += 1
                
                # Linha em branco entre filiais
                row_num += 1
            
            # Aba 2: Análise de Crescimento
            df_crescimento.to_excel(writer, sheet_name='Análise Crescimento', index=False)
            ws_crescimento = writer.sheets['Análise Crescimento']
            
            # Aba 3: Análise por Filial
            df_filial.to_excel(writer, sheet_name='Análise por Filial', index=False)
            ws_filial = writer.sheets['Análise por Filial']
            
            # Formata todas as abas
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Ajusta largura das colunas
            ws_detalhado.column_dimensions['A'].width = 12
            ws_detalhado.column_dimensions['B'].width = 50
            ws_detalhado.column_dimensions['C'].width = 20
            ws_detalhado.column_dimensions['D'].width = 18
            ws_detalhado.column_dimensions['E'].width = 20
            ws_detalhado.column_dimensions['F'].width = 18
            ws_detalhado.column_dimensions['G'].width = 20
            ws_detalhado.column_dimensions['H'].width = 18
            
            ws_detalhado.freeze_panes = 'A3'
            
            # Formata aba Análise Crescimento
            for cell in ws_crescimento[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
                ws_crescimento.column_dimensions[col].width = 18
            
            for row in ws_crescimento.iter_rows(min_row=2, max_row=ws_crescimento.max_row):
                for cell in row:
                    cell.border = border
                    if cell.column in [4, 5, 6]:  # Outubro, Novembro, Dezembro
                        cell.number_format = 'R$ #,##0.00'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    elif cell.column in [7, 8, 9]:  # Crescimentos
                        cell.number_format = '0.00"%"'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        if cell.value and isinstance(cell.value, (int, float)):
                            if cell.value >= 0:
                                cell.font = Font(color="006100")
                            else:
                                cell.font = Font(color="C00000")
                    elif cell.column == 1:  # Código
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    elif cell.column == 3:  # Filial
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.font = Font(bold=True)
            
            ws_crescimento.freeze_panes = 'A2'
            
            # Formata aba Análise por Filial
            for cell in ws_filial[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
                ws_filial.column_dimensions[col].width = 18
            
            for row in ws_filial.iter_rows(min_row=2, max_row=ws_filial.max_row):
                for cell in row:
                    cell.border = border
                    if cell.column == 1:  # Filial
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.font = Font(bold=True)
                    elif cell.column in [2, 3, 4]:  # Outubro, Novembro, Dezembro
                        cell.number_format = 'R$ #,##0.00'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    elif cell.column in [5, 6, 7]:  # Crescimentos
                        cell.number_format = '0.00"%"'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        if cell.value and isinstance(cell.value, (int, float)):
                            if cell.value >= 0:
                                cell.font = Font(color="006100")
                            else:
                                cell.font = Font(color="C00000")
            
            ws_filial.freeze_panes = 'A2'
        
        print(f"💾 Dados salvos em Excel: {arquivo_excel}")
        print()
        
        # Mostra resumo no console
        print("=" * 100)
        print("📊 RESUMO DA ANÁLISE")
        print("=" * 100)
        print()
        
        # Agrupa por produto para mostrar resumo
        produtos_mostrados = set()
        for _, row in df_excel.iterrows():
            codprod = row['CODPROD']
            produto = row['PRODUTO']
            filial = row['FILIAL']
            
            if codprod not in produtos_mostrados:
                print(f"📦 {codprod} - {produto}")
                produtos_mostrados.add(codprod)
            
            print(f"   Filial {filial}:")
            print(f"      Outubro:   Valor = R$ {row['OUTUBRO_VALOR']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') + 
                  f" | Margem = {row['OUTUBRO_MARGEM']:.2f}%")
            print(f"      Novembro:  Valor = R$ {row['NOVEMBRO_VALOR']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') + 
                  f" | Margem = {row['NOVEMBRO_MARGEM']:.2f}%")
            print(f"      Dezembro:  Valor = R$ {row['DEZEMBRO_VALOR']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') + 
                  f" | Margem = {row['DEZEMBRO_MARGEM']:.2f}%")
            print()
        
        print("=" * 100)
        print("✅ ANÁLISE CONCLUÍDA!")
        print("=" * 100)
        
    except Exception as e:
        print(f"❌ Erro ao executar análise: {e}")
        import traceback
        traceback.print_exc()
    finally:
        conn.close()
        print("\n🔌 Conexão fechada.")

# portfolio-commit-ready: analise_produtos_1615_543_3meses
